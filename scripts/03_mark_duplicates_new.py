#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把資料夾內所有 Word 檔（.docx）的內容，逐一跟「豆皮丸資料庫」
（例如你合併出來的 merged.docx）比對，找出重複的文字段落，
直接在原始檔案裡用「刪除線」標註出來。

v5 更新（新增「標註出處」功能）：
- 讀取資料庫時會記住每一段文字屬於哪一份來源（例如 11丸Av），
  比對到重複內容時，會在刪除線後面用灰色小字標上出處，例如：
      這是一段重複的文字〔11丸Av〕
- 出處直接對應 merged.docx 裡的標題「檔案：11丸Av.jpg」，
  副檔名會自動去掉，標成〔11丸Av〕。
  （這正是 OCR 程式 doc.add_heading(f"檔案：{filename}") 產生的格式）
- 「檔案：xxx.jpg」這行標題本身不會列入比對，避免誤判成重複內容。
- 可先用 --list-sources 檢查「程式到底抓到哪些出處」，確認沒問題再正式跑。
- --annotate-mode 可選擇標在每個重複處旁邊（inline，預設），
  或整段結束後統一列出（end）。
- --no-annotate 則完全不標出處，行為跟 v4 一樣。

v4 既有功能：
- 會自動往下搜尋所有子資料夾（不只是你指定的那一層）。
- 遇到 .doc（舊版 Word 格式）、.txt、.rtf 檔案，
  會先用 macOS 內建的 textutil 工具自動轉成同名的 .docx
  （原始 .doc/.txt/.rtf 檔案完全不會被更動或搬移，
  只是在同一個資料夾多產生一份轉檔後的 .docx），
  接著才對這份新產生的 .docx 進行比對、標註刪除線。
- 沿用 v3 的空白與標點正規化。

v6 更新：
- 預設不再產生 .bak 備份檔，直接在原檔上標註。
  （需要備份時再自己加 --backup）

=== 事前準備 ===
pip3 install --upgrade python-docx
（textutil 是 macOS 內建工具，不用另外安裝）

=== 使用方式 ===
先確認出處抓得對不對（不會更動任何檔案）：
python3 mark_duplicates.py "/Users/c-lab/Desktop/佩峰/要標註的資料夾" \
    --database "/Users/c-lab/AI工具/merged.docx" --list-sources

確認後正式跑：
python3 mark_duplicates.py "/Users/c-lab/Desktop/佩峰/要標註的資料夾" \
    --database "/Users/c-lab/AI工具/merged.docx"

參數：
  folder           要處理的資料夾（會自動往下搜尋所有子資料夾）
  --database       豆皮丸資料庫的合併文件路徑（必填）
  --min-match      判定為重複的最小字元長度（忽略空格後計算，預設 6）
  --backup         備份原始 .docx 檔案成 .bak（預設不備份，直接原地改）
  --no-recursive   只掃描指定的這一層資料夾，不搜尋子資料夾
  --list-sources   只列出從資料庫辨識到的出處清單，不做任何比對與修改
  --no-annotate    不標註出處，只畫刪除線
  --annotate-mode  inline（標在每個重複處旁邊，預設）或 end（整段最後統一列出）
  --source-pattern 自訂出處的辨識規則（正規表示式，預設抓「檔案：xxx.jpg」開頭的標題行）
  --heading-source 額外把 Word 標題樣式的段落也當成出處
  --keep-source-line  連「檔案：xxx.jpg」這行也列入比對
"""

import os
import re
import shutil
import argparse
import subprocess
from bisect import bisect_right
from difflib import SequenceMatcher

from docx import Document
from docx.shared import Pt, RGBColor

PUNCT_MAP = str.maketrans({
    "：": ":",
    "，": ",",
    "、": ",",
    "（": "(",
    "）": ")",
    "　": " ",
    "；": ";",
    "！": "!",
    "？": "?",
})

CONVERTIBLE_EXTS = (".doc", ".txt", ".rtf")

# 預設的出處寫法：對應 OCR 程式產生的標題「檔案：11丸Av.jpg」
DEFAULT_SOURCE_PATTERN = r"^\s*檔案\s*[：:]\s*(\S.*?)\s*$"

# 出處名稱結尾若是這些副檔名就自動去掉（讓 11丸Av.jpg 變成 11丸Av）
IMAGE_EXTS = (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp", ".heic")

# 出處行最長幾個字才會被當成「標題行」（太長的通常是內文，不是出處）
MAX_SOURCE_LINE_LEN = 60

UNKNOWN_SOURCE = "未標示出處"


def build_normalized(text):
    return "".join(ch.translate(PUNCT_MAP) for ch in text if not ch.isspace())


def build_normalized_with_map(text):
    chars = []
    idx_map = []
    for i, ch in enumerate(text):
        if ch.isspace():
            continue
        chars.append(ch.translate(PUNCT_MAP))
        idx_map.append(i)
    return "".join(chars), idx_map


def is_heading_paragraph(paragraph):
    """這個段落是不是用了 Word 的標題樣式（合併時加上去的，不是原始內容）。"""
    try:
        style_name = paragraph.style.name or ""
    except Exception:
        return False
    lowered = style_name.lower()
    return lowered.startswith("heading") or lowered.startswith("title") or "標題" in style_name


def clean_label(label):
    """把「11丸Av.jpg」這種名稱去掉副檔名，變成「11丸Av」。"""
    label = label.strip()
    lowered = label.lower()
    for ext in IMAGE_EXTS:
        if lowered.endswith(ext):
            return label[: -len(ext)].strip()
    return label


def detect_source_label(paragraph, regex, use_heading):
    """判斷這個段落是不是一個「出處標題」，是的話回傳出處名稱，否則回傳 None。"""
    text = paragraph.text.strip()
    if not text or len(text) > MAX_SOURCE_LINE_LEN:
        return None

    if regex is not None:
        m = regex.search(text)
        if m:
            label = m.group(1) if m.groups() else m.group(0)
            label = clean_label(label)
            if label:
                return label

    if use_heading and is_heading_paragraph(paragraph):
        return clean_label(text)

    return None


def load_reference(database_path, source_pattern, use_heading, keep_source_line=False):
    """
    讀取資料庫，回傳：
      reference_normalized：整份資料庫正規化後的字串
      segments：[(起始位置, 結束位置, 出處名稱), ...]，位置對應上面那個字串
    """
    doc = Document(database_path)
    regex = re.compile(source_pattern) if source_pattern else None

    parts = []
    segments = []
    total = 0
    current_label = UNKNOWN_SOURCE
    current_start = 0

    for p in doc.paragraphs:
        label = detect_source_label(p, regex, use_heading)
        if label:
            if label != current_label:
                if total > current_start:
                    segments.append((current_start, total, current_label))
                current_label = label
                current_start = total
            if not keep_source_line:
                # 「檔案：xxx.jpg」這行是合併時加的標題，不是原始內容，不列入比對
                continue
        elif is_heading_paragraph(p) and not keep_source_line:
            # 例如 OCR 程式加的大標「豆皮丸辨識結果」，同樣不列入比對
            continue

        norm = build_normalized(p.text)
        if norm:
            parts.append(norm)
            total += len(norm)

    if total > current_start:
        segments.append((current_start, total, current_label))

    return "".join(parts), segments


def labels_for_range(segments, seg_starts, start, end):
    """找出資料庫中 [start, end) 這段文字屬於哪些出處。"""
    if not segments:
        return []
    i = bisect_right(seg_starts, start) - 1
    if i < 0:
        i = 0
    labels = []
    while i < len(segments) and segments[i][0] < end:
        seg_start, seg_end, label = segments[i]
        if seg_end > start and label not in labels:
            labels.append(label)
        i += 1
    return labels


def find_duplicate_spans(text, reference_normalized, min_match, segments, seg_starts):
    """回傳 [(起始, 結束, [出處, ...]), ...]，位置對應原始段落文字。"""
    if not text.strip():
        return []
    text_normalized, idx_map = build_normalized_with_map(text)
    if not text_normalized:
        return []

    sm = SequenceMatcher(None, text_normalized, reference_normalized)
    raw_spans = []
    for block in sm.get_matching_blocks():
        if block.size >= min_match:
            start_orig = idx_map[block.a]
            end_orig = idx_map[block.a + block.size - 1] + 1
            labels = labels_for_range(segments, seg_starts, block.b, block.b + block.size)
            raw_spans.append((start_orig, end_orig, labels))

    raw_spans.sort(key=lambda s: (s[0], s[1]))
    merged = []
    for start, end, labels in raw_spans:
        if merged and start <= merged[-1][1]:
            prev_start, prev_end, prev_labels = merged[-1]
            combined = list(prev_labels)
            for label in labels:
                if label not in combined:
                    combined.append(label)
            merged[-1] = (prev_start, max(prev_end, end), combined)
        else:
            merged.append((start, end, list(labels)))
    return merged


def add_source_run(paragraph, labels):
    """在段落中加上一段灰色小字的出處註記。"""
    if not labels:
        return
    run = paragraph.add_run("〔" + "、".join(labels) + "〕")
    run.font.strike = False
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    run.font.italic = True


def rebuild_paragraph_with_strike(paragraph, spans, annotate, annotate_mode):
    text = paragraph.text
    if not spans:
        return
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)

    pos = 0
    all_labels = []
    for start, end, labels in spans:
        if start > pos:
            paragraph.add_run(text[pos:start])
        strike_run = paragraph.add_run(text[start:end])
        strike_run.font.strike = True
        for label in labels:
            if label not in all_labels:
                all_labels.append(label)
        if annotate and annotate_mode == "inline":
            add_source_run(paragraph, labels)
        pos = end
    if pos < len(text):
        paragraph.add_run(text[pos:])

    if annotate and annotate_mode == "end":
        add_source_run(paragraph, all_labels)


def process_file(filepath, reference_normalized, min_match, backup,
                 segments, seg_starts, annotate, annotate_mode):
    doc = Document(filepath)
    changed = False
    found_labels = []
    for p in doc.paragraphs:
        spans = find_duplicate_spans(p.text, reference_normalized, min_match, segments, seg_starts)
        if spans:
            rebuild_paragraph_with_strike(p, spans, annotate, annotate_mode)
            changed = True
            for _, _, labels in spans:
                for label in labels:
                    if label not in found_labels:
                        found_labels.append(label)
    if changed:
        if backup:
            shutil.copy2(filepath, filepath + ".bak")
        doc.save(filepath)
    return changed, found_labels


def find_target_files(root_folder, recursive):
    """找出所有要處理的檔案：直接是 .docx 的，或需要先轉檔的 .doc/.txt/.rtf。"""
    docx_files = []
    convertible_files = []

    if recursive:
        walker = os.walk(root_folder)
    else:
        walker = [(root_folder, [], os.listdir(root_folder))]

    for dirpath, _, filenames in walker:
        for name in filenames:
            if name.startswith("~$") or name.startswith("."):
                continue
            full_path = os.path.join(dirpath, name)
            ext = os.path.splitext(name)[1].lower()
            if ext == ".docx":
                docx_files.append(full_path)
            elif ext in CONVERTIBLE_EXTS:
                convertible_files.append(full_path)

    return sorted(docx_files), sorted(convertible_files)


def convert_to_docx(filepath):
    """用 macOS 內建 textutil 把 .doc/.txt/.rtf 轉成同資料夾的 .docx，回傳新檔案路徑。"""
    base, _ = os.path.splitext(filepath)
    output_path = base + ".docx"
    if os.path.exists(output_path):
        # 已經轉過，直接沿用，不重複轉檔
        return output_path
    subprocess.run(
        ["textutil", "-convert", "docx", filepath, "-output", output_path],
        check=True,
        capture_output=True,
    )
    return output_path


def main():
    parser = argparse.ArgumentParser(description="比對 Word 文件與豆皮丸資料庫，標註重複內容為刪除線並註明出處")
    parser.add_argument("folder", help="要比對並標註的資料夾")
    parser.add_argument("--database", required=True, help="豆皮丸資料庫合併文件路徑（例如 merged.docx）")
    parser.add_argument("--min-match", type=int, default=6, help="判定為重複的最小字元長度（忽略空格後計算，預設 6）")
    parser.add_argument("--backup", action="store_true", help="備份原始 .docx 檔案成 .bak（預設不備份）")
    parser.add_argument("--no-recursive", action="store_true", help="只掃描指定的這一層資料夾，不搜尋子資料夾")
    parser.add_argument("--no-annotate", action="store_true", help="不標註出處，只畫刪除線")
    parser.add_argument("--annotate-mode", choices=["inline", "end"], default="inline",
                        help="出處標在每個重複處旁邊（inline，預設）或整段最後統一列出（end）")
    parser.add_argument("--source-pattern", default=DEFAULT_SOURCE_PATTERN,
                        help="辨識出處的正規表示式（預設抓「檔案：xxx.jpg」這種標題）")
    parser.add_argument("--heading-source", action="store_true",
                        help="額外把所有 Word 標題樣式的段落也當成出處（預設只認「檔案：」開頭那種）")
    parser.add_argument("--keep-source-line", action="store_true",
                        help="把「檔案：xxx.jpg」這行標題本身也列入比對（預設不列入）")
    parser.add_argument("--list-sources", action="store_true",
                        help="只列出從資料庫辨識到的出處清單，不做任何比對與修改")
    args = parser.parse_args()

    annotate = not args.no_annotate

    print("讀取豆皮丸資料庫...")
    reference_normalized, segments = load_reference(
        args.database,
        source_pattern=args.source_pattern,
        use_heading=args.heading_source,
        keep_source_line=args.keep_source_line,
    )
    seg_starts = [s[0] for s in segments]
    print(f"資料庫共 {len(reference_normalized)} 字（已忽略空白）")

    unique_labels = []
    for _, _, label in segments:
        if label not in unique_labels:
            unique_labels.append(label)
    print(f"辨識到 {len(unique_labels)} 個出處：{'、'.join(unique_labels) if unique_labels else '（無）'}\n")

    if args.list_sources:
        print("=" * 50)
        print("以下是程式從資料庫抓到的出處與各自的字數：")
        for start, end, label in segments:
            print(f"  {label}：{end - start} 字")
        print("=" * 50)
        print("如果抓錯或抓不到，請用 --source-pattern 指定正確的規則，")
        print("確認無誤後再拿掉 --list-sources 正式執行。")
        return

    if annotate and unique_labels == [UNKNOWN_SOURCE]:
        print("⚠️  注意：資料庫裡沒有辨識到任何出處標題，所有重複處都只會標成「未標示出處」。")
        print("    建議先用 --list-sources 檢查，或用 --source-pattern 調整辨識規則。\n")

    docx_files, convertible_files = find_target_files(args.folder, recursive=not args.no_recursive)

    if convertible_files:
        print(f"找到 {len(convertible_files)} 份需要先轉檔的 .doc/.txt/.rtf 檔案，開始轉換...\n")
        for filepath in convertible_files:
            filename = os.path.basename(filepath)
            print(f"轉檔中: {filename} ...", end="", flush=True)
            try:
                new_docx = convert_to_docx(filepath)
                docx_files.append(new_docx)
                print(" ✅ 已轉成 .docx")
            except Exception as e:
                print(f" ❌ 轉檔失敗: {e}")
        docx_files = sorted(set(docx_files))
        print()

    if not docx_files:
        print(f"在 '{args.folder}'（含子資料夾）中沒有找到任何可處理的文件。")
        return

    print(f"共 {len(docx_files)} 份 .docx 文件，開始比對...\n")
    for idx, filepath in enumerate(docx_files, 1):
        rel_name = os.path.relpath(filepath, args.folder)
        print(f"[{idx}/{len(docx_files)}] 比對中: {rel_name} ...", end="", flush=True)
        try:
            changed, found_labels = process_file(
                filepath, reference_normalized, args.min_match,
                backup=args.backup,
                segments=segments, seg_starts=seg_starts,
                annotate=annotate, annotate_mode=args.annotate_mode,
            )
            if changed:
                if found_labels:
                    print(f" ✅ 已標註重複處（出處：{'、'.join(found_labels)}）")
                else:
                    print(" ✅ 已標註重複處")
            else:
                print(" 無重複")
        except Exception as e:
            print(f" ❌ 出錯: {e}")

    print("\n" + "=" * 50)
    print("完成！所有 .docx 檔案已在原地更新（重複內容以刪除線標註）。")
    if annotate:
        print("重複處後面的灰色小字〔◯◯〕就是該段文字在資料庫中的出處。")
    if convertible_files:
        print("原本的 .doc/.txt/.rtf 檔案完全沒被更動，轉檔後的 .docx 是新產生在同一個資料夾裡的檔案。")
    if args.backup:
        print("每份有異動的 .docx 檔案旁都留了一份 .bak 備份。")
    print("=" * 50)


if __name__ == "__main__":
    main()
