#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
批次使用 Google Gemini (gemini-2.5-flash-lite) 辨識橫豎混排圖片
包含 429 / 503 自動重試機制、安全過濾保護，並隱藏警告訊息

=== 事前準備 ===
API 金鑰改用環境變數讀取，不寫死在程式碼裡。
執行前，先在終端機設定（每次開新的終端機視窗都要重設一次）：

    export GEMINI_API_KEY="你的金鑰"

或者更方便的做法：把上面那行加進 ~/.zshrc 檔案最後一行，
存檔後重開一次終端機，之後就不用每次手動設定。
"""
import os
import glob
import time
import warnings
from PIL import Image
from google import genai
from docx import Document

# 隱藏警告提示訊息
warnings.filterwarnings("ignore")

# --------------------------------------------------
API_KEY = os.environ.get("GEMINI_API_KEY")
IMAGE_FOLDER = "/Users/c-lab/Desktop/佩峰/豆皮丸jpg檔(300小)"
OUTPUT_FOLDER = "/Users/c-lab/AI工具"
OUTPUT_FILENAME = "豆皮丸辨識結果_merged.docx"
MODEL_NAME = "gemini-2.5-flash-lite"
# --------------------------------------------------


def call_gemini_with_retry(client, prompt, image, max_retries=5):
    """遇流量限制或伺服器繁忙時自動重試"""
    for attempt in range(1, max_retries + 1):
        try:
            response = client.models.generate_content(
                model=MODEL_NAME, contents=[prompt, image]
            )
            # 如果被安全機制擋掉，response.candidates 會是空的，
            # 這時候直接讀 response.text 會噴錯，先檢查再讀取
            if not response.candidates:
                return "[無法辨識：內容被安全機制阻擋]"
            if not response.text:
                return "[無法辨識：回傳內容為空]"
            return response.text.strip()
        except Exception as e:
            if any(err in str(e) for err in ["503", "429", "RESOURCE_EXHAUSTED", "UNAVAILABLE"]):
                wait_time = attempt * 5
                print(f" (頻率限制/伺服器繁忙，等待 {wait_time} 秒後重試第 {attempt} 次...)", end="", flush=True)
                time.sleep(wait_time)
            else:
                raise e
    raise Exception("達到最大重試次數。")


def process_images():
    if not API_KEY:
        print("❌ 找不到 API 金鑰！請先在終端機執行：export GEMINI_API_KEY=\"你的金鑰\"")
        return

    os.makedirs(OUTPUT_FOLDER, exist_ok=True)
    output_docx_path = os.path.join(OUTPUT_FOLDER, OUTPUT_FILENAME)

    client = genai.Client(api_key=API_KEY)

    extensions = ('*.jpg', '*.jpeg', '*.png', '*.JPG', '*.JPEG', '*.PNG')
    image_paths = []
    for ext in extensions:
        image_paths.extend(glob.glob(os.path.join(IMAGE_FOLDER, ext)))
    image_paths = sorted(image_paths)

    if not image_paths:
        print(f"❌ 在 '{IMAGE_FOLDER}' 資料夾中找不到任何圖片！")
        return

    print(f"找到 {len(image_paths)} 張圖片，開始進行橫豎混排 OCR 辨識...\n")

    doc = Document()
    doc.add_heading('豆皮丸辨識結果', level=1)

    prompt = """
    這是一張繁體中文文件圖片，版面中可能包含橫排與豎排文字。
    請精準辨識圖片中的所有文字，並遵循以下規則：
    1. 按照中文正確的閱讀順序（區分標題、欄位、內文）整理輸出。
    2. 豎排文字請自動整理為正常的橫向閱讀文字段落。
    3. 完整提取所有文字，不要遺漏，也不要加入你自己的說明或評語。
    """

    for idx, img_path in enumerate(image_paths, 1):
        filename = os.path.basename(img_path)
        print(
            f"[{idx}/{len(image_paths)}] 辨識中: {filename} ...",
            end="",
            flush=True,
        )
        try:
            image = Image.open(img_path)
            result_text = call_gemini_with_retry(client, prompt, image)
            doc.add_heading(f"檔案：{filename}", level=2)
            for line in result_text.split("\n"):
                if line.strip():
                    doc.add_paragraph(line)
                else:
                    doc.add_paragraph("")
            print(" ✅ 完成")
            time.sleep(2)
        except Exception as e:
            print(f" ❌ 出錯: {e}")
            doc.add_heading(f"檔案：{filename}", level=2)
            doc.add_paragraph(f"[辨識失敗: {e}]")

    doc.save(output_docx_path)
    print("\n" + "=" * 50)
    print(f"所有圖片辨識完成！Word 檔已儲存至：\n{output_docx_path}")
    print("=" * 50)


if __name__ == "__main__":
    process_images()
